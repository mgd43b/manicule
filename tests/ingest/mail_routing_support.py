"""A pipeline over the *real* parser registry, and messages shaped the way real clients write them.

``tests/connectors/gmail_support.mail_pipeline`` resolves every media type to the same two
parsers, so a test built on it cannot say anything about routing. This one resolves chains the
way production does (``Container.parser_chain_names``) and runs the registered parsers, which is
what a mail attachment actually meets. Every address is at ``example.org``.
"""

from __future__ import annotations

import base64
import io
import tempfile
from collections.abc import Callable, Sequence
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

import docx
import openpyxl

from manicule.config.settings import Settings
from manicule.container import keys
from manicule.container.container import build_container
from manicule.core.fingerprints import ParseFingerprint
from manicule.ingest.middleware import MiddlewareRunner
from manicule.ingest.pipeline import IngestPipeline
from manicule.ingest.workers import InProcessRunner
from manicule.parsers.versions import parse_fingerprint
from manicule.plugins.registry import ComponentKind
from tests.fakes import HashEmbedder
from tests.ingest import fakes


class RealRoutes:
    """``routes[media_type]``, answered by the production resolver."""

    def __init__(self, resolve: object) -> None:
        self._resolve = resolve

    def __getitem__(self, media_type: str) -> Sequence[str]:
        return list(self._resolve(media_type))  # type: ignore[operator]


async def real_pipeline(
    *,
    store: fakes.MemoryIngestStore | None = None,
    blobs: fakes.MemoryBlobs | None = None,
    parse_fingerprints: Callable[[str], ParseFingerprint | None] = parse_fingerprint,
) -> tuple[IngestPipeline, fakes.MemoryIngestStore]:
    """The ingest pipeline with every shipped parser behind the production chain resolver."""
    directory = tempfile.mkdtemp()
    container = build_container(Settings(data_dir=Path(directory), cache_dir=Path(directory)))
    parsers = {
        record.name: await container.aget(keys.PARSER.named(record.name))
        for record in container.registry.records(ComponentKind.PARSER)
    }
    store = store or fakes.MemoryIngestStore()
    chunker = fakes.BlockChunker()
    pipeline = IngestPipeline(
        store=store,
        chunker=chunker,
        embedder=HashEmbedder(),
        vectors=fakes.MemoryVectors(),
        runner=InProcessRunner(parsers),
        resolve_chain=container.parser_chain_names,
        middleware=MiddlewareRunner([]),
        chunk_fingerprint=chunker.fingerprint,
        blobs=blobs,
        parse_fingerprints=parse_fingerprints,
    )
    return pipeline, store


def docx_bytes(text: str) -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(text: str) -> bytes:
    book = openpyxl.Workbook()
    book.active["A1"] = text  # type: ignore[index]
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def message(*parts: tuple[str, bytes, str, str | None]) -> bytes:
    """A message from ``(content-type, bytes, transfer-encoding, filename)`` attachments.

    Written by hand rather than by ``EmailMessage.add_attachment`` because the point is the
    Content-Type a client chose, including the ones the standard library would never emit
    (``application/msword`` for a ``.docx``, ``message/rfc822`` in base64).
    """
    out = [
        "From: pat@example.org\r\nTo: lee@example.org\r\nSubject: attachments\r\n"
        "Message-ID: <m@example.org>\r\nMIME-Version: 1.0\r\n"
        "Content-Type: multipart/mixed; boundary=BOUND\r\n\r\n"
        "--BOUND\r\nContent-Type: text/plain\r\n\r\nSee attached.\r\n"
    ]
    for content_type, data, encoding, filename in parts:
        body = (
            base64.encodebytes(data).decode().replace("\n", "\r\n")
            if encoding == "base64"
            else data.decode("latin-1")
        )
        disposition = (
            f'Content-Disposition: attachment; filename="{filename}"\r\n' if filename else ""
        )
        out.append(
            f"--BOUND\r\nContent-Type: {content_type}\r\n"
            f"Content-Transfer-Encoding: {encoding}\r\n{disposition}\r\n{body}\r\n"
        )
    out.append("--BOUND--\r\n")
    return "".join(out).encode("latin-1")


def enclosed_message(subject: str, body: str) -> bytes:
    inner = EmailMessage()
    inner["From"] = "sam@example.org"
    inner["Subject"] = subject
    inner.set_content(body + "\n")
    return inner.as_bytes(policy=SMTP)
